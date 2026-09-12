"""建置步驟：從官方 F01/F02/F03 範本抽出「可協填欄位目錄」→ data/rag/form_fields.json。

欄位標籤／說明是**範本用語**，不得入庫：本腳本唯讀 `data/original/`，輸出到 `data/rag/`
（gitignored，同 retrieval_map.json 政策）。座標來源：
  - F01：review_config.yaml `f01.main_sheet` 第 2 列標籤、第 3 列說明、欄 A..T，資料列 4。
  - F02：`AI系統固有風險分級評估表` 第 2–46 列；A 欄命中 f02_scoring.yaml `questions` 的題號，
        或加成因子列（`uplift_factors.*.row`）；答案格 C{row}。題組成員的 description 取母題題幹。
  - F03：review_config.yaml `f03.checklist` 座標：A 欄項次（`^\\d+-\\d+`）、C 管理議題、D 檢查項目，
        目標格為佐證說明欄 I{row}；顯示標籤用 prompts.catalog.short_name_for_item。

`query_terms` 是本檔手寫的小字典（通用治理用語，不是範本句子），用來補檢索的詞面重疊。

用法：uv run python scripts/extract_form_fields.py   （只印計數，不印任何範本文字）
"""

from __future__ import annotations

import json
import re
import sys
import warnings
from pathlib import Path

import openpyxl
import yaml
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from govcheck.prompts.catalog import short_name_for_item  # noqa: E402

warnings.filterwarnings("ignore")

ORIGINAL = ROOT / "data" / "original"
SRC_F01 = ORIGINAL / "附件一：AI-R02-F01 AI系統資訊表.xlsx"
SRC_F02 = ORIGINAL / "附件二：AI-R02-F02 AI風險評鑑.xlsm"
SRC_F03 = ORIGINAL / "附件三：AI-R02-F03 AI系統提案規劃暨上線檢核表.xlsx"
REVIEW_CONFIG = ROOT / "src" / "govcheck" / "config" / "review_config.yaml"
SCORING_CONFIG = ROOT / "src" / "govcheck" / "config" / "f02_scoring.yaml"
OUT_JSON = ROOT / "data" / "rag" / "form_fields.json"

SCHEMA_VERSION = 1

F01_LAST_COL = "T"
F02_SHEET = "AI系統固有風險分級評估表"
F02_ROWS = range(2, 47)
F02_ANSWER_COL = "C"
F02_FACTOR_LABEL_RE = re.compile(r"^加成因子\s*([123])$")

# F01 M..T 欄 review_config 沒有語意鍵（審查規則用不到），這裡補上讓 field_id 穩定可讀。
F01_EXTRA_KEYS = {
    "M": "ai_user_unit",
    "N": "ai_program_provider",
    "O": "ai_model_provider",
    "P": "ai_platform_provider",
    "Q": "managing_unit",
    "R": "business_name",
    "S": "affected_parties",
    "T": "other_stakeholders",
}

# 檢索補充詞（通用用語）。key = field_id。
QUERY_TERMS: dict[str, list[str]] = {
    # F01
    "filing_unit": ["填表單位", "申請單位", "提案單位", "部門", "處", "科"],
    "filer": ["填表人", "承辦人", "聯絡人", "窗口"],
    "filing_date": ["日期", "填表日期", "民國", "年", "月", "日"],
    "project_name": ["專案名稱", "服務名稱", "系統名稱", "計畫名稱"],
    "ai_usage": ["AI 應用", "用途", "功能", "場景", "使用方式"],
    "affected_system_name": ["系統名稱", "介接系統", "影響系統"],
    "affected_system_code": ["系統代碼", "代號", "系統編號"],
    "system_owner": ["System Owner", "系統擁有者", "系統負責人"],
    "ap_owner": ["AP Owner", "應用程式負責人", "AP 負責單位"],
    "ai_app_type": ["應用類型", "生成式", "預測", "分類", "推薦", "對話"],
    "business_goal": ["業務目標", "目的", "效益", "預期成效", "KPI"],
    "tech_type": ["技術類型", "模型", "LLM", "機器學習", "深度學習", "RAG"],
    "ai_user_unit": ["使用單位", "使用者", "業務單位"],
    "ai_program_provider": ["程式提供", "開發單位", "廠商", "供應商"],
    "ai_model_provider": ["模型提供", "模型來源", "廠商", "基礎模型"],
    "ai_platform_provider": ["平台", "運作環境", "雲端", "地端", "IT 單位"],
    "managing_unit": ["管理單位", "主管單位", "權責單位"],
    "business_name": ["業務名稱", "業務別", "適用業務"],
    "affected_parties": ["受影響對象", "客戶", "員工", "使用者", "群體"],
    "other_stakeholders": ["利害關係人", "相關單位", "合作夥伴"],
    # F02（通用治理用語）
    "UC-01": ["客戶權益", "重大影響", "客戶"],
    "UC-02": ["員工", "考績", "人員權利", "就業"],
    "UC-03": ["監理", "法規遵循", "主管機關"],
    "UC-04-01": ["自動決策", "最終決策", "無人工"],
    "UC-04-02": ["告警", "異常", "人類介入", "HOTL"],
    "UC-04-03": ["人類決策", "人工確認", "HITL", "建議"],
    "UC-04-04": ["內容產生", "生成", "草稿", "不涉及決策"],
    "UC-05-01": ["無法申覆", "不可修改"],
    "UC-05-02": ["申訴", "重新審視", "調整結果"],
    "UC-05-03": ["申訴", "人工重新處理", "救濟"],
    "UC-05-04": ["不涉及決策"],
    "UC-06-01": ["中斷", "無替代方案"],
    "UC-06-02": ["中斷", "人員替代", "備援"],
    "UC-06-03": ["人為中止", "替代方案", "備援", "關閉"],
    "UC-07": ["告知", "揭露", "AI 技術", "互動", "透明"],
    "D-01-01": ["個人資料", "客戶個資", "直接使用"],
    "D-01-02": ["去識別化", "匿名化", "假名化"],
    "D-01-03": ["資料來源", "可能包含", "個資"],
    "D-01-04": ["不使用個資", "無個人資料"],
    "D-02-01": ["內部資料", "行內資料", "公司資料"],
    "D-02-02": ["公開資料", "網路", "爬取"],
    "D-02-03": ["購買", "外部資料", "非公開"],
    "D-02-04": ["購買", "外部", "公開資料"],
    "D-03": ["機敏資料", "內部機密", "營業秘密"],
    "D-04": ["隱私", "資料保護", "加密", "存取控制", "資安"],
    "M-01": ["內部開發", "自行開發", "採購", "外購"],
    "M-02": ["模型", "內部開發", "第三方模型", "開源"],
    "M-03-01": ["內部維運", "自行維運"],
    "M-03-02": ["委外", "第三方", "SaaS", "維運廠商"],
    "factor1": ["非重大", "不具業務重大性", "內部工具", "輔助"],
    "factor2": ["個人工作", "效率", "日常", "輔助工具"],
    "factor3": ["涵蓋範圍", "50%", "受影響對象", "廣度"],
    # F03（依 catalog 短名衍生）
    "1-1": ["適法性", "法規", "法遵", "合規評估"],
    "1-2": ["決策說明", "可解釋", "解釋性", "說明機制"],
    "1-3": ["測試", "驗測", "驗證結果", "準確率"],
    "1-4": ["溝通", "救濟", "申訴", "客訴"],
    "1-5": ["人工監督", "人為介入", "覆核"],
    "1-6": ["定期檢視", "簽核", "審查週期"],
    "2-1": ["業務目標", "用途", "預期產出", "效益"],
    "2-2": ["委外", "廠商", "供應商", "合約"],
    "3-1": ["資料來源", "代表性", "樣本"],
    "3-2": ["資料蒐集", "偏差", "偏誤"],
    "3-3": ["個資", "隱私", "去識別化", "保護措施"],
    "3-4": ["資料取得", "傳輸", "儲存", "加密"],
    "4-1": ["個人屬性", "公平性", "敏感屬性", "歧視"],
    "4-2": ["監控", "模型監控", "告警"],
    "4-3": ["KPI", "偏移", "漂移", "監測"],
    "4-4": ["版本", "紀錄", "留存", "模型登錄"],
    "4-5": ["模型挑選", "超參數", "調校"],
    "5-1": ["變更管理", "變更流程", "版本控管"],
    "5-2": ["權限", "帳號", "存取控制"],
    "5-3": ["資訊安全", "存取程序", "資安控管"],
}


def _text(value: object) -> str:
    return " ".join(str(value).split()) if value is not None else ""


def extract_f01(review_cfg: dict) -> list[dict]:
    cfg = review_cfg["f01"]
    columns: dict[str, str] = {**cfg["columns"], **F01_EXTRA_KEYS}
    required = {row["col"] for row in cfg["required_columns"]}
    wb = openpyxl.load_workbook(SRC_F01, data_only=True, read_only=True)
    ws = wb[cfg["main_sheet"]]
    data_row = int(cfg.get("data_start_row", 4))
    fields = []
    for index in range(1, openpyxl.utils.column_index_from_string(F01_LAST_COL) + 1):
        col = get_column_letter(index)
        label = _text(ws.cell(row=2, column=index).value)
        if not label:
            continue
        field_id = columns.get(col, f"col_{col.lower()}")
        fields.append({
            "form": "f01",
            "field_id": field_id,
            "label": label,
            "description": _text(ws.cell(row=3, column=index).value),
            "cell": f"{col}{data_row}",
            "kind": "text",
            "required": col in required,
            "query_terms": QUERY_TERMS.get(field_id, []),
        })
    return fields


def extract_f02(scoring_cfg: dict) -> list[dict]:
    questions = scoring_cfg["questions"]
    groups = scoring_cfg["groups"]
    group_kind_of: dict[str, str] = {}
    for kind in ("single_choice", "multi_choice"):
        for group_id in groups.get(kind, {}):
            group_kind_of[group_id] = kind
    factor_rows = {
        int(spec["row"]): factor_id
        for factor_id, spec in scoring_cfg.get("uplift_factors", {}).items()
    }
    wb = openpyxl.load_workbook(SRC_F02, data_only=True, read_only=True)
    ws = wb[F02_SHEET]
    stems: dict[str, str] = {}
    fields = []
    for row in F02_ROWS:
        raw_id = _text(ws.cell(row=row, column=1).value)
        text = _text(ws.cell(row=row, column=2).value)
        if not raw_id:
            continue
        if raw_id in group_kind_of:
            stems[raw_id] = text  # 題組母題：不是答案列，題幹留給成員當 description
            continue
        factor_match = F02_FACTOR_LABEL_RE.match(raw_id)
        if factor_match:
            field_id = factor_rows.get(row) or f"factor{factor_match.group(1)}"
            group = None
        elif raw_id in questions:
            field_id = raw_id
            group = raw_id.rsplit("-", 1)[0] if raw_id.count("-") == 2 else None
            if group not in group_kind_of:
                group = None
        else:
            continue
        fields.append({
            "form": "f02",
            "field_id": field_id,
            "label": text,
            "description": stems.get(group, "") if group else "",
            "cell": f"{F02_ANSWER_COL}{row}",
            "kind": "yes_no",
            "group": group,
            "group_kind": group_kind_of.get(group) if group else None,
            "required": True,
            "query_terms": QUERY_TERMS.get(field_id, []),
        })
    return fields


def extract_f03(review_cfg: dict) -> list[dict]:
    cfg = review_cfg["f03"]
    checklist = cfg["checklist"]
    pattern = re.compile(checklist["item_id_pattern"])
    wb = openpyxl.load_workbook(SRC_F03, data_only=True, read_only=True)
    ws = wb[cfg["sheet"]]
    id_col = openpyxl.utils.column_index_from_string(checklist["item_id_col"])
    topic_col = openpyxl.utils.column_index_from_string(checklist["topic_col"])
    desc_col = openpyxl.utils.column_index_from_string(checklist["item_desc_col"])
    target_col = checklist["evidence_proposal_col"]
    fields = []
    for row in range(int(checklist["data_start_row"]), ws.max_row + 1):
        item_id = _text(ws.cell(row=row, column=id_col).value)
        if not item_id or not pattern.match(item_id):
            continue
        topic = _text(ws.cell(row=row, column=topic_col).value)
        fields.append({
            "form": "f03",
            "field_id": item_id,
            "label": short_name_for_item(item_id) or topic or item_id,
            "description": _text(ws.cell(row=row, column=desc_col).value),
            "cell": f"{target_col}{row}",
            "kind": "text",
            "required": True,
            "query_terms": [*([topic] if topic else []), *QUERY_TERMS.get(item_id, [])],
        })
    return fields


def main() -> None:
    for source in (SRC_F01, SRC_F02, SRC_F03):
        if not source.exists():
            raise SystemExit(f"找不到官方範本：{source.name}（data/original 未連結？）")
    review_cfg = yaml.safe_load(REVIEW_CONFIG.read_text(encoding="utf-8"))
    scoring_cfg = yaml.safe_load(SCORING_CONFIG.read_text(encoding="utf-8"))
    f01 = extract_f01(review_cfg)
    f02 = extract_f02(scoring_cfg)
    f03 = extract_f03(review_cfg)
    payload = {
        "schema_version": SCHEMA_VERSION,
        "_source": "data/original 官方範本（由 scripts/extract_form_fields.py 自動產出，勿手改、勿入庫）",
        "fields": [*f01, *f02, *f03],
    }
    OUT_JSON.parent.mkdir(parents=True, exist_ok=True)
    OUT_JSON.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    f02_questions = sum(1 for field in f02 if not field["field_id"].startswith("factor"))
    f02_factors = len(f02) - f02_questions
    print(
        f"寫入 {OUT_JSON.relative_to(ROOT)}：F01 {len(f01)} 欄、"
        f"F02 {f02_questions} 題 + {f02_factors} 加成因子、F03 {len(f03)} 項"
    )


if __name__ == "__main__":
    main()
