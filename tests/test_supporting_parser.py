from __future__ import annotations

import io
import zipfile

import openpyxl
from docx import Document

from govcheck.models import DocumentParseStatus
from govcheck.parsers.supporting import ParserSettings, TsParserAdapter, load_parser_settings


def _settings(*formats: str, **overrides) -> ParserSettings:
    values = {
        "enabled_formats": set(formats),
        "max_file_bytes": 1024 * 1024,
        "max_pages": 20,
        "temp_root": None,
    }
    values.update(overrides)
    return ParserSettings(**values)


def test_text_source_parses_to_stable_artifact(tmp_path):
    source = tmp_path / "機密名稱.txt"
    source.write_text("公平性測試已完成，結果符合門檻。", encoding="utf-8")

    artifact, usage = TsParserAdapter(_settings(".txt")).parse(source)

    assert artifact.status is DocumentParseStatus.PARSED
    assert artifact.parser_name == "ts_parser"
    assert artifact.source_units[0].locator_value == "1"
    assert artifact.filename_hash != source.name
    assert usage.parser_vlm_calls == 0


def test_xlsx_preserves_sheet_locator(tmp_path):
    source = tmp_path / "evidence.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.title = "測試結果"
    workbook.active.append(["項目", "結果"])
    workbook.active.append(["公平性", "通過"])
    workbook.save(source)

    artifact, _ = TsParserAdapter(_settings(".xlsx")).parse(source)

    assert artifact.status is DocumentParseStatus.PARSED
    assert artifact.source_units[0].locator_type == "sheet"
    assert artifact.source_units[0].locator_value == "測試結果"


def test_invalid_pdf_is_unreadable_not_not_found(tmp_path):
    source = tmp_path / "evidence.pdf"
    source.write_bytes(b"not-a-pdf")

    artifact, _ = TsParserAdapter(_settings(".pdf")).parse(source)

    assert artifact.status is DocumentParseStatus.UNREADABLE


def test_vlm_format_fails_closed_without_approved_endpoint(tmp_path):
    source = tmp_path / "screen.png"
    source.write_bytes(b"\x89PNG\r\n\x1a\nsynthetic")

    artifact, usage = TsParserAdapter(_settings(".png", vlm_enabled=False)).parse(source)

    assert artifact.status is DocumentParseStatus.PARSER_UNAVAILABLE
    assert usage.parser_vlm_calls == 1


def test_image_ocr_usage_is_tracked_separately(tmp_path):
    source = tmp_path / "screen.png"
    source.write_bytes(b"\x89PNG\r\n\x1a\nsynthetic")

    artifact, usage = TsParserAdapter(_settings(
        ".png", ocr_enabled=True, model_host="https://unapproved.example"
    )).parse(source)

    assert artifact.status is DocumentParseStatus.PARSER_UNAVAILABLE
    assert usage.parser_ocr_calls == 1
    assert usage.parser_vlm_calls == 0


def test_docx_native_strategy_parses_text(tmp_path):
    source = tmp_path / "evidence.docx"
    document = Document()
    document.add_paragraph("模型卡已記錄測試範圍與結果。")
    document.save(source)

    artifact, usage = TsParserAdapter(_settings(".docx", word_strategy="native")).parse(source)

    assert artifact.status is DocumentParseStatus.PARSED
    assert "模型卡" in artifact.source_units[0].content
    assert usage.parser_vlm_calls == 0


def test_pptx_requires_enabled_approved_vlm(tmp_path):
    source = tmp_path / "evidence.pptx"
    with zipfile.ZipFile(source, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types />")

    artifact, usage = TsParserAdapter(_settings(".pptx")).parse(source)

    assert artifact.status is DocumentParseStatus.PARSER_UNAVAILABLE
    assert usage.parser_vlm_calls == 1


def test_msg_disabled_format_is_unsupported(tmp_path):
    source = tmp_path / "mail.msg"
    source.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1synthetic")

    artifact, _ = TsParserAdapter(_settings(".txt")).parse(source)

    assert artifact.status is DocumentParseStatus.UNSUPPORTED


def test_msg_enabled_but_dependency_missing_is_parser_unavailable(tmp_path):
    source = tmp_path / "mail.msg"
    source.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1synthetic")

    artifact, _ = TsParserAdapter(_settings(".msg")).parse(source)

    assert artifact.status is DocumentParseStatus.PARSER_UNAVAILABLE


def test_zip_bomb_limit_returns_unsupported(tmp_path):
    source = tmp_path / "large.docx"
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", "A" * 50_000)
    source.write_bytes(buffer.getvalue())

    artifact, _ = TsParserAdapter(_settings(
        ".docx", max_zip_uncompressed_bytes=100, max_zip_ratio=2
    )).parse(source)

    assert artifact.status is DocumentParseStatus.UNSUPPORTED
    assert artifact.warnings[0] in {"zip_uncompressed_limit_exceeded", "zip_ratio_limit_exceeded"}


def test_duplicate_document_is_counted_as_cache_hit(tmp_path):
    source = tmp_path / "same.txt"
    source.write_text("same", encoding="utf-8")

    artifacts, usage = TsParserAdapter(_settings(".txt")).parse_many([source, source])

    assert len(artifacts) == 1
    assert usage.cache_hits == 1


# ── 設定讀取韌性與可重現性 ───────────────────────────────────────────────────

def test_blank_env_values_fall_back_to_defaults(monkeypatch):
    """.env 慣例會留下 `KEY=` 空值；不可因此丟 ValueError 拖垮整份審查。"""
    for key in ("GOVCHECK_PARSER_TIMEOUT", "GOVCHECK_PARSER_MAX_FILE_BYTES",
                "GOVCHECK_PARSER_MAX_PAGES", "GOVCHECK_PARSER_WORD_STRATEGY"):
        monkeypatch.setenv(key, "")

    settings = load_parser_settings()

    assert settings.request_timeout == 60.0
    assert settings.max_pages == 200
    assert settings.word_strategy == "auto"


def test_garbage_env_value_falls_back_instead_of_raising(monkeypatch):
    monkeypatch.setenv("GOVCHECK_PARSER_MAX_PAGES", "not-a-number")
    assert load_parser_settings().max_pages == 200


def test_config_hash_is_stable_for_equal_settings():
    """parser_config_hash 是稽核用的設定指紋：同設定必須算出同一個值。"""
    a = ParserSettings(
        enabled_formats={".pdf", ".txt", ".docx"},
        approved_model_hosts={"localhost", "127.0.0.1"},
    )
    b = ParserSettings(
        enabled_formats={".docx", ".pdf", ".txt"},      # 同內容、不同插入順序
        approved_model_hosts={"127.0.0.1", "localhost"},
    )
    assert a.safe_hash() == b.safe_hash()
