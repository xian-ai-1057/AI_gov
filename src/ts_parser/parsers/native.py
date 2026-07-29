"""Local, deterministic parsers for text-bearing evidence documents."""

from __future__ import annotations

import csv
import io
from io import BytesIO
from typing import Any

import openpyxl
import pdfplumber
from docx import Document

from ts_parser.parsers.base_parser import BaseParser, ParserUnavailableError, UnsupportedFileError


def _cell(value: Any) -> str:
    text = "" if value is None else str(value)
    return text.replace("|", "\\|").replace("\n", " ").strip()


def _markdown(rows: list[list[Any]]) -> str:
    if not rows:
        return ""
    width = max(len(row) for row in rows)
    normalized = [row + [None] * (width - len(row)) for row in rows]
    header = [_cell(v) or f"欄位{i + 1}" for i, v in enumerate(normalized[0])]
    body = normalized[1:]
    lines = [f"| {' | '.join(header)} |", f"| {' | '.join(['---'] * width)} |"]
    lines.extend(f"| {' | '.join(_cell(v) for v in row)} |" for row in body)
    return "\n".join(lines)


class TextParser(BaseParser):
    def parse(self, file_bytes: BytesIO, filename: str, **kwargs: Any) -> list[dict[str, Any]]:
        raw = file_bytes.getvalue()
        text = None
        for encoding in ("utf-8-sig", "utf-8", "big5", "cp950"):
            try:
                text = raw.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        if text is None:
            raise ValueError("無法辨識文字檔編碼")
        ext = self.extension(filename)
        if ext in {".csv", ".tsv"}:
            rows = list(csv.reader(io.StringIO(text), delimiter="," if ext == ".csv" else "\t"))
            text = _markdown(rows)
        return [{"content": text, "page": "1"}]


class SpreadsheetParser(BaseParser):
    def parse(self, file_bytes: BytesIO, filename: str, **kwargs: Any) -> list[dict[str, Any]]:
        if self.extension(filename) == ".xls":
            raise ParserUnavailableError(".xls 需要尚未啟用的 xlrd parser")
        workbook = openpyxl.load_workbook(file_bytes, read_only=True, data_only=True, keep_links=False)
        try:
            chunks = []
            for sheet in workbook.worksheets:
                rows = [list(row) for row in sheet.iter_rows(values_only=True)]
                chunks.append({"content": _markdown(rows), "page": sheet.title})
            return chunks
        finally:
            workbook.close()


class DocxParser(BaseParser):
    def parse(self, file_bytes: BytesIO, filename: str, **kwargs: Any) -> list[dict[str, Any]]:
        if self.extension(filename) == ".doc":
            raise ParserUnavailableError(".doc 需要 Office conversion 或 Melody parser")
        document = Document(file_bytes)
        blocks = [p.text.strip() for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            blocks.append(_markdown([[cell.text for cell in row.cells] for row in table.rows]))
        return [{"content": "\n\n".join(blocks), "page": "1"}]


class PdfParser(BaseParser):
    def __init__(self, *, max_pages: int = 200) -> None:
        self.max_pages = max_pages

    def parse(self, file_bytes: BytesIO, filename: str, **kwargs: Any) -> list[dict[str, Any]]:
        chunks = []
        with pdfplumber.open(file_bytes) as pdf:
            if len(pdf.pages) > self.max_pages:
                raise UnsupportedFileError("PDF 頁數超過安全限制")
            for index, page in enumerate(pdf.pages, start=1):
                chunks.append({"content": page.extract_text() or "", "page": str(index)})
        return chunks


class MsgParser(BaseParser):
    def parse(self, file_bytes: BytesIO, filename: str, **kwargs: Any) -> list[dict[str, Any]]:
        raise ParserUnavailableError("MSG parser 需要 extract-msg；目前功能旗標預設關閉")
